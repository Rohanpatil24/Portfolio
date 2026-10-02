# core/resume_builder.py
import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .models import PersonalInfo, Experience, Project, Skill, Education, Certification

def get_portfolio_snapshot():
    info = PersonalInfo.objects.first()
    experiences = Experience.objects.all().order_by('order', '-id')
    projects = Project.objects.all().order_by('order', '-id')
    skills = Skill.objects.all()
    education = Education.objects.all()
    certifications = Certification.objects.all()
    return info, experiences, projects, skills, education, certifications

def generate_resume_docx():
    """Generates an ATS-compliant Word document (.docx)."""
    info, experiences, projects, skills, education, certifications = get_portfolio_snapshot()
    doc = Document()

    # Set 0.6 inch standard ATS margins
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # 1. Header (Name & Contact)
    name = info.name if info and info.name else "Rohan Patil"
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = p_name.add_run(name.upper())
    run_name.bold = True
    run_name.font.size = Pt(18)
    run_name.font.name = "Calibri"

    contacts = []
    if info:
        if info.location: contacts.append(info.location)
        if info.email: contacts.append(info.email)
        if info.phone: contacts.append(info.phone)
        if info.linkedin_url: contacts.append(f"LinkedIn: {info.linkedin_url}")
        if info.github_url: contacts.append(f"GitHub: {info.github_url}")

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(8)
    run_contact = p_contact.add_run(" | ".join(contacts))
    run_contact.font.size = Pt(9.5)
    run_contact.font.name = "Calibri"
    run_contact.font.color.rgb = RGBColor(70, 70, 70)

    def add_section_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(13, 148, 136) # Teal accent

    # 2. Professional Summary
    if info and info.about_text:
        add_section_heading("Professional Summary")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(info.about_text.strip())
        r.font.size = Pt(10)
        r.font.name = "Calibri"

    # 3. Core Competencies / Technical Skills
    if skills.exists():
        add_section_heading("Technical Skills")
        categories = {'backend': 'Backend & Databases', 'frontend': 'Frontend & Interface', 'tools': 'Cloud & DevOps'}
        for cat_key, cat_name in categories.items():
            cat_skills = [s.name for s in skills if s.category == cat_key]
            if cat_skills:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                r_title = p.add_run(f"• {cat_name}: ")
                r_title.bold = True
                r_title.font.size = Pt(9.5)
                r_title.font.name = "Calibri"
                r_skills = p.add_run(", ".join(cat_skills))
                r_skills.font.size = Pt(9.5)
                r_skills.font.name = "Calibri"

    # 4. Work Experience
    if experiences.exists():
        add_section_heading("Work Experience")
        for exp in experiences:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            r_role = p.add_run(f"{exp.role} | {exp.company}")
            r_role.bold = True
            r_role.font.size = Pt(10.5)
            r_role.font.name = "Calibri"
            r_dur = p.add_run(f" ({exp.duration})")
            r_dur.italic = True
            r_dur.font.size = Pt(9.5)
            r_dur.font.name = "Calibri"

            for bullet in exp.description.split('\n'):
                line = bullet.strip().lstrip('•-*').strip()
                if line:
                    bp = doc.add_paragraph(style='List Bullet')
                    bp.paragraph_format.space_after = Pt(1)
                    br = bp.add_run(line)
                    br.font.size = Pt(9.5)
                    br.font.name = "Calibri"

    # 5. Projects
    if projects.exists():
        add_section_heading("Key Projects")
        for proj in projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            r_title = p.add_run(f"{proj.title} ({proj.get_role_display()})")
            r_title.bold = True
            r_title.font.size = Pt(10)
            r_title.font.name = "Calibri"

            if proj.tech_stack:
                r_stack = p.add_run(f" — Tech Stack: {proj.tech_stack}")
                r_stack.italic = True
                r_stack.font.size = Pt(9)
                r_stack.font.name = "Calibri"

            desc_p = doc.add_paragraph()
            desc_p.paragraph_format.space_after = Pt(3)
            dr = desc_p.add_run(proj.description.strip())
            dr.font.size = Pt(9.5)
            dr.font.name = "Calibri"

    # 6. Education
    if education.exists():
        add_section_heading("Education")
        for edu in education:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r_deg = p.add_run(f"• {edu.title} — {edu.institution}")
            r_deg.bold = True
            r_deg.font.size = Pt(9.5)
            r_deg.font.name = "Calibri"
            r_date = p.add_run(f" ({edu.date_completed})")
            r_date.font.size = Pt(9)
            r_date.font.name = "Calibri"

    # 7. Licenses & Certifications
    if certifications.exists():
        add_section_heading("Certifications")
        for cert in certifications:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"• {cert.title} — {cert.institution} ({cert.date_completed})")
            r.font.size = Pt(9.5)
            r.font.name = "Calibri"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def generate_resume_pdf():
    """Generates an ATS-compliant PDF using standard fonts and linear hierarchy."""
    info, experiences, projects, skills, education, certifications = get_portfolio_snapshot()
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.55 * inch,
        rightMargin=0.55 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch
    )

    styles = getSampleStyleSheet()
    name_style = ParagraphStyle('NameStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=18, leading=22, alignment=1, textColor=colors.HexColor('#0f172a'))
    contact_style = ParagraphStyle('ContactStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, leading=12, alignment=1, textColor=colors.HexColor('#475569'))
    heading_style = ParagraphStyle('HeadingStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10.5, leading=14, spaceBefore=7, spaceAfter=2, textColor=colors.HexColor('#0d9488'))
    body_style = ParagraphStyle('BodyStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=12.5, textColor=colors.HexColor('#1e293b'))
    bullet_style = ParagraphStyle('BulletStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=8.8, leading=12, leftIndent=12, textColor=colors.HexColor('#1e293b'))

    story = []

    # 1. Header
    name = info.name if info and info.name else "Rohan Patil"
    story.append(Paragraph(name.upper(), name_style))
    story.append(Spacer(1, 3))

    contacts = []
    if info:
        if info.location: contacts.append(info.location)
        if info.email: contacts.append(info.email)
        if info.phone: contacts.append(info.phone)
        if info.linkedin_url: contacts.append(f"LinkedIn: {info.linkedin_url}")
        if info.github_url: contacts.append(f"GitHub: {info.github_url}")

    story.append(Paragraph(" | ".join(contacts), contact_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceBefore=3, spaceAfter=6))

    def add_pdf_section(title):
        story.append(Paragraph(title.upper(), heading_style))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#0d9488'), spaceBefore=1, spaceAfter=4))

    # 2. Professional Summary
    if info and info.about_text:
        add_pdf_section("Professional Summary")
        story.append(Paragraph(info.about_text.strip(), body_style))
        story.append(Spacer(1, 4))

    # 3. Technical Skills
    if skills.exists():
        add_pdf_section("Technical Skills")
        categories = {'backend': 'Backend & Databases', 'frontend': 'Frontend & Interface', 'tools': 'Cloud & DevOps'}
        for cat_key, cat_name in categories.items():
            cat_skills = [s.name for s in skills if s.category == cat_key]
            if cat_skills:
                story.append(Paragraph(f"<b>&bull; {cat_name}:</b> {', '.join(cat_skills)}", body_style))
        story.append(Spacer(1, 4))

    # 4. Work Experience
    if experiences.exists():
        add_pdf_section("Work Experience")
        for exp in experiences:
            exp_header = f"<b>{exp.role}</b> — {exp.company} <i>({exp.duration})</i>"
            story.append(Paragraph(exp_header, body_style))
            for bullet in exp.description.split('\n'):
                line = bullet.strip().lstrip('•-*').strip()
                if line:
                    story.append(Paragraph(f"&bull; {line}", bullet_style))
            story.append(Spacer(1, 3))

    # 5. Projects
    if projects.exists():
        add_pdf_section("Key Projects")
        for proj in projects:
            title_text = f"<b>{proj.title}</b> ({proj.get_role_display()}) &mdash; <i>{proj.tech_stack}</i>"
            story.append(Paragraph(title_text, body_style))
            story.append(Paragraph(proj.description.strip(), bullet_style))
            story.append(Spacer(1, 3))

    # 6. Education
    if education.exists():
        add_pdf_section("Education")
        for edu in education:
            story.append(Paragraph(f"<b>&bull; {edu.title}</b> — {edu.institution} ({edu.date_completed})", body_style))
        story.append(Spacer(1, 3))

    # 7. Certifications
    if certifications.exists():
        add_pdf_section("Certifications")
        for cert in certifications:
            story.append(Paragraph(f"&bull; {cert.title} — {cert.institution} ({cert.date_completed})", body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer